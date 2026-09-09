# -*- coding: utf-8 -*-
"""네이버 검색광고(NSA) × GA4 키워드 리포트 결합기

사용법
    python merge_nsa_ga.py                  같은 폴더의 CSV 를 자동 인식
    python merge_nsa_ga.py --scope 전체     파워링크 외 유형(쇼핑검색 등)도 포함
    python merge_nsa_ga.py --out 결과.xlsx  출력 파일명 지정

입력 — 이 스크립트와 같은 폴더에 두면 됩니다
    NSA*.csv    네이버 검색광고 키워드 리포트 (캠페인유형 · 캠페인 · 광고그룹 · 키워드 …)
    GA*.csv     GA4 리포트 (세션 캠페인 · 세션 수동 검색어 · 기기 카테고리 …)

결합 방식
    키 = 날짜 × 브랜드 × 키워드 × 기기(PC/MO)
      브랜드  NSA 캠페인명 NAV_{MI|IT|EBM}_SA_… ↔ GA 세션 캠페인 {mi|it|ebm}_paidsearch
      기기    NSA 캠페인명 접미사 _pc/_mo ↔ GA 기기 카테고리(desktop=PC, mobile·tablet=MO)
    브랜드 + 기기가 캠페인을 유일하게 결정하므로 캠페인 단위 안분이 필요 없다.
    광고그룹만 일부 조합이 둘로 갈리며, 이때만 GA 값을 클릭 비중으로 안분한다.

출력 시트
    매칭 상세 (광고그룹)   광고그룹까지 펼친 결합 결과. GA 배분 = 실측 / 안분
    미매칭 요약           매칭 현황 + 미매칭 사유별 집계
    미매칭 Raw            한쪽에만 있는 조합 전체
"""
import argparse
import csv
import datetime as dt
import glob
import io
import os
import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:                                    # 콘솔이 cp949 여도 한글이 깨지지 않게
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

DEVMAP = {'desktop': 'PC', 'mobile': 'MO', 'tablet': 'MO'}
CAMP_RE = re.compile(r'NAV_([A-Za-z]+)_SA_(?:pf_)?[a-z]*_?(pc|mo)(?:_|$)', re.I)
PERIOD_RE = re.compile(r'(\d{4})\.(\d{2})\.\d{2}\.~')

INK, AMBER, GRAYL = 'FF1A1A1A', 'FFFFBE32', 'FFF2F2F2'
THIN = Side(style='thin', color='FFD9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [('날짜', 11, 'd'), ('브랜드', 8, 'c'), ('캠페인', 24, 's'),
        ('광고그룹', 26, 's'), ('키워드', 20, 's'), ('기기', 6, 'c'),
        ('GA 배분', 10, 'c'),
        ('노출수', 11, 'n'), ('클릭수', 10, 'n'), ('총비용', 12, 'n'),
        ('CTR', 9, 'p'), ('CPC', 9, 'n'),
        ('NSA 전환수', 11, 'n'), ('NSA 전환매출', 14, 'n'),
        ('세션수', 10, 'n'), ('구매', 8, 'n'), ('구매 수익', 14, 'n'),
        ('세션/클릭', 10, 'p'), ('ROAS', 9, 'p0')]
UNCOLS = [c if c[0] != 'GA 배분' else ('구분', 9, 'c') for c in COLS[:17]]


def num(v):
    try:
        return float(str(v).replace(',', '').strip() or 0)
    except ValueError:
        return 0.0


def norm(s):
    return ''.join(str(s).split()).lower()


def dv(a, b, m=1.0):
    return (a / b * m) if b else 0.0


def d8(s):
    """2026.08.01. / 20260801 → 20260801"""
    return re.sub(r'[^0-9]', '', str(s))[:8]


def to_date(s):
    return dt.date(int(s[:4]), int(s[4:6]), int(s[6:8]))


def find(pattern, folder, label):
    hits = [f for f in glob.glob(os.path.join(folder, pattern))
            if not os.path.basename(f).startswith('~$')]
    if not hits:
        sys.exit('[중단] %s 파일을 찾지 못했습니다 (%s)' % (label, pattern))
    if len(hits) > 1:
        sys.exit('[중단] %s 파일이 여러 개입니다: %s'
                 % (label, ', '.join(os.path.basename(h) for h in hits)))
    return hits[0]


# ── 읽기 ───────────────────────────────────────────────
def load_nsa(path, scope):
    rows = list(csv.reader(io.open(path, encoding='utf-8-sig')))
    period = ''
    m = PERIOD_RE.search(rows[0][0] if rows and rows[0] else '')
    if m:
        period = '%s-%s' % (m.group(1), m.group(2))
    hdr = rows[1]
    H = {k: i for i, k in enumerate(hdr)}
    need = ('일별', '캠페인유형', '캠페인', '광고그룹', '키워드', '총비용',
            '노출수', '클릭수', '구매완료 전환수', '구매완료 전환매출액(원)')
    miss = [k for k in need if k not in H]
    if miss:
        sys.exit('[중단] NSA 리포트에 없는 컬럼: %s' % ', '.join(miss))

    out = defaultdict(lambda: defaultdict(lambda: [0.0] * 5))
    label, types, skipped = {}, defaultdict(int), 0
    for r in rows[2:]:
        if not r or len(r) < len(hdr):
            continue
        t = r[H['캠페인유형']]
        types[t] += 1
        if scope and t not in scope:
            continue
        mm = CAMP_RE.match(r[H['캠페인']])
        if not mm:
            skipped += 1
            continue
        k = (d8(r[H['일별']]), mm.group(1).upper(), norm(r[H['키워드']]),
             mm.group(2).upper())
        label.setdefault(k, r[H['키워드']].strip())
        a = out[k][(r[H['캠페인']], r[H['광고그룹']])]
        for i, v in enumerate((num(r[H['총비용']]), num(r[H['노출수']]),
                               num(r[H['클릭수']]), num(r[H['구매완료 전환수']]),
                               num(r[H['구매완료 전환매출액(원)']]))):
            a[i] += v
    return out, label, period, dict(types), skipped


def load_ga(path):
    lines = io.open(path, encoding='utf-8-sig').read().splitlines()
    try:
        st = next(i for i, l in enumerate(lines) if l.startswith('세션 소스/매체'))
    except StopIteration:
        sys.exit('[중단] GA 리포트에서 헤더행(세션 소스/매체 …)을 찾지 못했습니다')
    rows = list(csv.reader(lines[st:]))
    G = {k: i for i, k in enumerate(rows[0])}
    need = ('날짜', '세션 캠페인', '세션 수동 검색어', '기기 카테고리', '구매',
            '구매 수익', '세션수')
    miss = [k for k in need if k not in G]
    if miss:
        sys.exit('[중단] GA 리포트에 없는 컬럼: %s' % ', '.join(miss))

    out = defaultdict(lambda: [0.0] * 3)
    for r in rows[1:]:
        if not r or len(r) < len(rows[0]) or not r[0].strip():
            continue                                   # 총합계 · 빈 행 제외
        gc = r[G['세션 캠페인']].strip()
        k = (d8(r[G['날짜']]), gc.split('_')[0].upper() if gc else '?',
             norm(r[G['세션 수동 검색어']]),
             DEVMAP.get(r[G['기기 카테고리']].strip().lower(), '?'))
        a = out[k]
        for i, v in enumerate((num(r[G['세션수']]), num(r[G['구매']]),
                               num(r[G['구매 수익']]))):
            a[i] += v
    return out


# ── 쓰기 ───────────────────────────────────────────────
def style(ws, cols, nrow):
    for j, (nm, w, _k) in enumerate(cols, 1):
        c = ws.cell(1, j)
        c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=INK)
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), nrow)
    fmt = {'n': '#,##0', 'p': '0.0%', 'p0': '0%', 'd': 'yyyy-mm-dd'}
    for i in range(2, nrow + 1):
        for j, (_nm, _w, kind) in enumerate(cols, 1):
            c = ws.cell(i, j)
            c.font = Font(name='맑은 고딕', size=9.5)
            c.border = BORDER
            if kind in fmt:
                c.number_format = fmt[kind]
                c.alignment = Alignment(
                    horizontal='center' if kind == 'd' else 'right')
            elif kind == 'c':
                c.alignment = Alignment(horizontal='center')
            else:
                c.alignment = Alignment(horizontal='left')
            if i % 2 == 0:
                c.fill = PatternFill('solid', fgColor=GRAYL)


def summary_sheet(ws, blocks):
    ws.column_dimensions['A'].width = 40
    for j, w in zip('BCDEFGH', (10, 12, 11, 13, 11, 10, 14)):
        ws.column_dimensions[j].width = w
    r = 1
    for title, rowset in blocks:
        c = ws.cell(r, 1, title)
        c.font = Font(name='맑은 고딕', size=11, bold=True)
        r += 1
        head = rowset[0]
        pc = len(head) if head[-1].endswith('비중') else None
        for j, v in enumerate(head, 1):
            c = ws.cell(r, j, v)
            c.font = Font(name='맑은 고딕', size=9.5, bold=True, color='FFFFFFFF')
            c.fill = PatternFill('solid', fgColor=INK)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        r += 1
        for i, row in enumerate(rowset[1:]):
            tot = str(row[0]) in ('합계', '소계')
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v)
                c.font = Font(name='맑은 고딕', size=9.5, bold=tot)
                c.border = BORDER
                if j == 1:
                    c.alignment = Alignment(horizontal='left')
                else:
                    c.number_format = '0.0%' if (pc and j == pc) else '#,##0'
                    c.alignment = Alignment(horizontal='right')
                if tot:
                    c.fill = PatternFill('solid', fgColor=AMBER)
                elif i % 2 == 1:
                    c.fill = PatternFill('solid', fgColor=GRAYL)
            r += 1
        r += 2


def main():
    ap = argparse.ArgumentParser(description='NSA × GA4 키워드 리포트 결합')
    ap.add_argument('--dir', default=os.path.dirname(os.path.abspath(__file__)),
                    help='CSV 가 있는 폴더 (기본: 스크립트 위치)')
    ap.add_argument('--scope', default='파워링크',
                    help="NSA 캠페인유형. '전체' 로 두면 모든 유형 포함")
    ap.add_argument('--out', default=None, help='출력 xlsx 경로')
    a = ap.parse_args()

    scope = None if a.scope == '전체' else tuple(a.scope.split(','))
    nsa_csv = find('NSA*.csv', a.dir, 'NSA 리포트')
    ga_csv = find('GA*.csv', a.dir, 'GA 리포트')
    print('입력  NSA : %s' % os.path.basename(nsa_csv))
    print('      GA  : %s' % os.path.basename(ga_csv))

    nsa, label, period, types, skipped = load_nsa(nsa_csv, scope)
    ga = load_ga(ga_csv)
    print('기간  %s   범위 %s' % (period or '(미상)', a.scope))
    print('      NSA 유형 %s' % ' / '.join('%s %d행' % (k, v)
                                          for k, v in types.items()))
    if skipped:
        print('      ! 캠페인명에서 브랜드·기기를 못 읽어 제외한 행 %d' % skipped)

    keys = sorted(set(nsa) & set(ga))
    rows, split = [], 0
    for k in keys:
        parts = nsa[k]
        camps = set(c for c, _g in parts)
        if len(camps) > 1:
            sys.exit('[중단] 캠페인이 여러 개인 조합: %s %s' % (k, camps))
        ca = camps.pop()
        cost, imp, clk = [sum(p[i] for p in parts.values()) for i in range(3)]
        ses, pur, rev = ga[k]
        n = len(parts)
        if n > 1:
            split += 1
        base = clk if clk else (imp if imp else 0)
        for (_c, agn), p in sorted(parts.items(), key=lambda x: -x[1][2]):
            if n == 1:
                w = 1.0
            elif base:
                w = (p[2] if clk else p[1]) / base
            else:
                w = 1.0 / n
            rows.append([to_date(k[0]), k[1], ca, agn, label[k], k[3],
                         '실측' if n == 1 else '안분',
                         p[1], p[2], p[0], dv(p[2], p[1]), dv(p[0], p[2]),
                         p[3], p[4], ses * w, pur * w, rev * w,
                         dv(ses * w, p[2]), dv(rev * w, p[0])])
    rows.sort(key=lambda r: -r[8])

    un = []
    for k in sorted(set(nsa) - set(ga)):
        for (ca, agn), p in sorted(nsa[k].items(), key=lambda x: -x[1][2]):
            un.append([to_date(k[0]), k[1], ca, agn, label[k], k[3], 'NSA만',
                       p[1], p[2], p[0], dv(p[2], p[1]), dv(p[0], p[2]),
                       p[3], p[4], 0, 0, 0])
    for k in sorted(set(ga) - set(nsa)):
        s, pu, gr = ga[k]
        un.append([to_date(k[0]), k[1], '-', '-', k[2], k[3], 'GA만',
                   0, 0, 0, 0, 0, 0, 0, s, pu, gr])
    un.sort(key=lambda r: (-r[8], -r[14]))
    nonly = [r for r in un if r[6] == 'NSA만']
    gonly = [r for r in un if r[6] == 'GA만']

    def sm(rs, idx):
        return [sum(r[i] for r in rs) for i in idx]

    IDX = (7, 8, 9, 14, 15, 16)
    b1 = [['구분', '조합 수', '노출수', '클릭수', '총비용', '세션수', '구매',
           '구매 수익'],
          ['매칭 (양쪽)', len(rows)] + sm(rows, IDX),
          ['NSA만 (GA 세션 없음)', len(nonly)] + sm(nonly, IDX),
          ['GA만 (NSA 클릭 없음)', len(gonly)] + sm(gonly, IDX),
          ['합계', len(rows) + len(un)] + sm(rows + un, IDX)]

    def band(r):
        if r[4].strip() == '-':
            return '키워드 미특정 ("-")'
        return ('클릭 10회 이상' if r[8] >= 10 else
                '클릭 1~9회' if r[8] >= 1 else '클릭 0회 (노출만 발생)')

    b2 = [['NSA만 — 사유', '조합 수', '노출수', '클릭수', '총비용', '클릭 비중']]
    ncl = sum(r[8] for r in nonly) or 1
    for lb in ('키워드 미특정 ("-")', '클릭 10회 이상', '클릭 1~9회',
               '클릭 0회 (노출만 발생)'):
        rs = [r for r in nonly if band(r) == lb]
        v = sm(rs, (7, 8, 9))
        b2.append([lb, len(rs)] + v + [v[1] / ncl])
    b2.append(['소계', len(nonly)] + sm(nonly, (7, 8, 9)) + [1.0])

    nkw = set(k[2] for k in nsa)

    def why(r):
        k = norm(r[4])
        if k in ('(notset)', '(notprovided)'):
            return '검색어 미수집 ((not set) · (not provided))'
        if 'season_fall' in k:
            return '캠페인명이 검색어로 기록 (season_fall 등)'
        return ('같은 키워드가 다른 브랜드에만 존재' if k in nkw
                else 'NSA 대상 유형에 없는 검색어')

    b3 = [['GA만 — 사유', '조합 수', '세션수', '구매', '구매 수익', '세션 비중']]
    gse = sum(r[14] for r in gonly) or 1
    seen = {}
    for r in gonly:
        seen.setdefault(why(r), []).append(r)
    for lb in sorted(seen, key=lambda x: -sum(r[14] for r in seen[x])):
        v = sm(seen[lb], (14, 15, 16))
        b3.append([lb, len(seen[lb])] + v + [v[0] / gse])
    b3.append(['소계', len(gonly)] + sm(gonly, (14, 15, 16)) + [1.0])

    # 합계 검증 — 결합 전후 총량이 보존됐는지
    src_clk = sum(p[2] for v in nsa.values() for p in v.values())
    src_ses = sum(v[0] for v in ga.values())
    got_clk = sum(r[8] for r in rows + un)
    got_ses = sum(r[14] for r in rows + un)
    ok = abs(src_clk - got_clk) < 1 and abs(src_ses - got_ses) < 1
    print('검증  클릭 %s / 세션 %s  → %s'
          % (format(got_clk, ',.0f'), format(got_ses, ',.0f'),
             '원본과 일치' if ok else '!! 불일치'))
    if not ok:
        sys.exit('[중단] 합계가 원본과 맞지 않습니다')

    out = a.out or os.path.join(a.dir, 'NSA_GA_매칭_%s.xlsx' % (period or 'out'))
    wb = openpyxl.Workbook()
    ws = wb.create_sheet('매칭 상세 (광고그룹)')
    ws.append([c[0] for c in COLS])
    for r in rows:
        ws.append(r)
    style(ws, COLS, len(rows) + 1)

    summary_sheet(wb.create_sheet('미매칭 요약'),
                  [('① 매칭 현황', b1), ('② NSA만 — 왜 GA 세션이 없나', b2),
                   ('③ GA만 — 왜 NSA 클릭이 없나', b3)])

    ws = wb.create_sheet('미매칭 Raw')
    ws.append([c[0] for c in UNCOLS])
    for r in un:
        ws.append(r)
    style(ws, UNCOLS, len(un) + 1)

    wb.remove(wb['Sheet'])
    try:
        wb.save(out)
    except PermissionError:
        sys.exit('[중단] 파일이 열려 있습니다. 닫고 다시 실행하세요: %s'
                 % os.path.basename(out))
    print('결과  매칭 %d행 (안분 대상 %d조합) / 미매칭 %d행'
          % (len(rows), split, len(un)))
    print('      %s' % out)


if __name__ == '__main__':
    main()
