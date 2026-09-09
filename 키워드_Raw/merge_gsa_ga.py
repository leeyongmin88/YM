# -*- coding: utf-8 -*-
"""구글 검색광고(GSA) × GA4 리포트 결합 — 날짜 · 캠페인 · 광고그룹 · 검색어

사용법
    python merge_gsa_ga.py                     같은 폴더의 CSV 를 자동 인식
    python merge_gsa_ga.py --min-clicks 0      미매칭 Raw 에 클릭 0 행까지 포함
    python merge_gsa_ga.py --out 결과.xlsx     출력 파일명 지정

입력 — 이 스크립트와 같은 폴더에 두면 됩니다
    GSA_키워드*.csv   구글 광고 검색어 리포트 (UTF-16 · 탭 구분)
    GSA_GA*.csv       GA4 리포트 (UTF-8 · 쉼표 구분)

결합 방식
    키 = 날짜 × 캠페인 × 광고그룹 × 검색어  (네 값 모두 양쪽 리포트에 있음)
    날짜는 숫자만 남겨 맞추고(2026-08-01 = 20260801), 광고그룹 · 검색어는
    앞뒤 공백 정리 + 대소문자 무시로 비교한다. 안분 없이 전부 실측이다.

출력 시트
    매칭            네 키가 모두 맞는 조합
    미매칭 요약      매칭 현황 + 미매칭 사유별 집계
    미매칭 Raw      한쪽에만 있는 조합 (기본: 클릭 1회 이상만)
"""
import argparse
import csv
import datetime as dt
import glob
import io
import os
import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

INK, AMBER, GRAYL = 'FF1A1A1A', 'FFFFBE32', 'FFF2F2F2'
THIN = Side(style='thin', color='FFD9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = [('날짜', 11, 'd'), ('캠페인', 30, 's'), ('광고그룹', 26, 's'),
        ('검색어', 28, 's'),
        ('비용', 11, 'n'), ('노출수', 10, 'n'), ('클릭수', 9, 'n'),
        ('CTR', 9, 'p'), ('CPC', 9, 'n'),
        ('세션수', 9, 'n'), ('구매', 8, 'n'), ('구매 수익', 13, 'n'),
        ('세션/클릭', 10, 'p'), ('구매전환율', 11, 'p2'), ('ROAS', 9, 'p0')]
UNCOLS = COLS[:4] + [('구분', 9, 'c')] + COLS[4:12]


def num(v):
    try:
        return float(str(v).replace(',', '').replace('"', '').strip() or 0)
    except ValueError:
        return 0.0


def norm(s):
    return ' '.join(str(s).split()).lower()


def d8(s):
    return re.sub(r'[^0-9]', '', str(s))[:8]


def dv(a, b, m=1.0):
    return (a / b * m) if b else 0.0


def find(pattern, folder, label):
    hits = [f for f in glob.glob(os.path.join(folder, pattern))
            if not os.path.basename(f).startswith('~$')]
    if not hits:
        sys.exit('[중단] %s 파일을 찾지 못했습니다 (%s)' % (label, pattern))
    if len(hits) > 1:
        sys.exit('[중단] %s 파일이 여러 개입니다: %s'
                 % (label, ', '.join(os.path.basename(h) for h in hits)))
    return hits[0]


def read_text(path):
    """구글 리포트는 UTF-16, GA 는 UTF-8 로 내려온다"""
    raw = open(path, 'rb').read()
    for enc in ('utf-16', 'utf-8-sig', 'cp949'):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    sys.exit('[중단] 인코딩을 알 수 없습니다: %s' % os.path.basename(path))


# ── 읽기 ───────────────────────────────────────────────
def load_ad(path):
    lines = read_text(path).splitlines()
    try:
        hi = next(i for i, l in enumerate(lines) if l.split('\t')[0].strip() == '일')
    except StopIteration:
        sys.exit('[중단] 광고 리포트에서 헤더행(일 / 캠페인 / …)을 찾지 못했습니다')
    period = lines[1].strip() if hi >= 2 else ''
    rows = list(csv.reader(lines[hi:], delimiter='\t'))
    H = {k.strip(): i for i, k in enumerate(rows[0])}
    need = ('일', '캠페인', '광고그룹', '검색어', '비용', '노출수', '클릭수')
    miss = [k for k in need if k not in H]
    if miss:
        sys.exit('[중단] 광고 리포트에 없는 컬럼: %s' % ', '.join(miss))

    out = defaultdict(lambda: [0.0, 0.0, 0.0])
    label, camps = {}, defaultdict(lambda: [0.0, 0.0, 0.0])
    for r in rows[1:]:
        if not r or len(r) < len(rows[0]) or not r[0].strip():
            continue
        k = (d8(r[H['일']]), r[H['캠페인']].strip(),
             norm(r[H['광고그룹']]), norm(r[H['검색어']]))
        label.setdefault(k, (r[H['광고그룹']].strip(), r[H['검색어']].strip()))
        rec = (num(r[H['비용']]), num(r[H['노출수']]), num(r[H['클릭수']]))
        a, c = out[k], camps[k[1]]
        for i, v in enumerate(rec):
            a[i] += v
            c[i] += v
    return out, label, period, camps


def load_ga(path):
    lines = read_text(path).splitlines()
    try:
        hi = next(i for i, l in enumerate(lines) if l.startswith('날짜,'))
    except StopIteration:
        sys.exit('[중단] GA 리포트에서 헤더행(날짜, …)을 찾지 못했습니다')
    rows = list(csv.reader(lines[hi:]))
    G = {k.strip(): i for i, k in enumerate(rows[0])}
    need = ('날짜', '세션 Google Ads 캠페인', '세션 Google Ads 광고그룹 이름',
            '세션 Google Ads 검색어', '세션수', '구매', '구매 수익')
    miss = [k for k in need if k not in G]
    if miss:
        sys.exit('[중단] GA 리포트에 없는 컬럼: %s' % ', '.join(miss))

    out = defaultdict(lambda: [0.0, 0.0, 0.0])
    label = {}
    for r in rows[1:]:
        if not r or len(r) < len(rows[0]) or not r[0].strip():
            continue                                  # 총합계 · 빈 행 제외
        k = (d8(r[G['날짜']]), r[G['세션 Google Ads 캠페인']].strip(),
             norm(r[G['세션 Google Ads 광고그룹 이름']]),
             norm(r[G['세션 Google Ads 검색어']]))
        label.setdefault(k, (r[G['세션 Google Ads 광고그룹 이름']].strip(),
                             r[G['세션 Google Ads 검색어']].strip()))
        a = out[k]
        for i, v in enumerate((num(r[G['세션수']]), num(r[G['구매']]),
                               num(r[G['구매 수익']]))):
            a[i] += v
    return out, label


# ── 쓰기 ───────────────────────────────────────────────
def style(ws, cols, nrow):
    for j, (nm, w, _k) in enumerate(cols, 1):
        c = ws.cell(1, j)
        c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=INK)
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), nrow)
    fmt = {'n': '#,##0', 'p': '0.0%', 'p2': '0.00%', 'p0': '0%',
           'd': 'yyyy-mm-dd'}
    for i in range(2, nrow + 1):
        for j, (_nm, _w, kind) in enumerate(cols, 1):
            c = ws.cell(i, j)
            c.font = Font(name='맑은 고딕', size=9.5)
            c.border = BORDER
            if kind in fmt:
                c.number_format = fmt[kind]
                c.alignment = Alignment(
                    horizontal='center' if kind == 'd' else 'right')
            elif kind == 'c':
                c.alignment = Alignment(horizontal='center')
            else:
                c.alignment = Alignment(horizontal='left')
            if i % 2 == 0:
                c.fill = PatternFill('solid', fgColor=GRAYL)


def summary_sheet(ws, blocks):
    ws.column_dimensions['A'].width = 40
    for j, w in zip('BCDEFGH', (10, 12, 11, 10, 11, 10, 14)):
        ws.column_dimensions[j].width = w
    r = 1
    for title, rowset in blocks:
        c = ws.cell(r, 1, title)
        c.font = Font(name='맑은 고딕', size=11, bold=True)
        r += 1
        head = rowset[0]
        pc = len(head) if head[-1].endswith('비중') else None
        for j, v in enumerate(head, 1):
            c = ws.cell(r, j, v)
            c.font = Font(name='맑은 고딕', size=9.5, bold=True, color='FFFFFFFF')
            c.fill = PatternFill('solid', fgColor=INK)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        r += 1
        for i, row in enumerate(rowset[1:]):
            tot = str(row[0]) in ('합계', '소계')
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v)
                c.font = Font(name='맑은 고딕', size=9.5, bold=tot)
                c.border = BORDER
                if j == 1:
                    c.alignment = Alignment(horizontal='left')
                elif isinstance(v, str):
                    c.alignment = Alignment(horizontal='center')
                else:
                    c.number_format = '0.0%' if (pc and j == pc) else '#,##0'
                    c.alignment = Alignment(horizontal='right')
                if tot:
                    c.fill = PatternFill('solid', fgColor=AMBER)
                elif i % 2 == 1:
                    c.fill = PatternFill('solid', fgColor=GRAYL)
            r += 1
        r += 2


def to_date(s):
    return dt.date(int(s[:4]), int(s[4:6]), int(s[6:8]))


def main():
    ap = argparse.ArgumentParser(description='GSA × GA4 리포트 결합')
    ap.add_argument('--dir', default=os.path.dirname(os.path.abspath(__file__)))
    ap.add_argument('--min-clicks', type=int, default=1,
                    help='미매칭 Raw 에 실을 최소 클릭 수 (기본 1)')
    ap.add_argument('--out', default=None)
    a = ap.parse_args()

    ad_csv = find('GSA_키워드*.csv', a.dir, '광고 리포트')
    ga_csv = find('GSA_GA*.csv', a.dir, 'GA 리포트')
    print('입력  광고: %s' % os.path.basename(ad_csv))
    print('      GA  : %s' % os.path.basename(ga_csv))

    ad, alab, period, camps = load_ad(ad_csv)
    ga, glab = load_ga(ga_csv)
    gcamps = set(k[1] for k in ga)
    print('기간  %s' % (period or '(미상)'))
    print('      광고 %d조합 / 캠페인 %d개 (GA 추적 %d개)'
          % (len(ad), len(camps), len(gcamps)))

    keys = set(ad) & set(ga)
    rows = []
    for k in sorted(keys):
        cost, imp, clk = ad[k]
        ses, pur, rev = ga[k]
        agn, kw = alab.get(k) or glab.get(k)
        rows.append([to_date(k[0]), k[1], agn, kw, cost, imp, clk,
                     dv(clk, imp), dv(cost, clk), ses, pur, rev,
                     dv(ses, clk), dv(pur, ses), dv(rev, cost)])
    rows.sort(key=lambda r: (-r[6], -r[9]))

    un = []
    for k in sorted(set(ad) - keys):
        cost, imp, clk = ad[k]
        agn, kw = alab[k]
        un.append([to_date(k[0]), k[1], agn, kw, '광고만',
                   cost, imp, clk, dv(clk, imp), dv(cost, clk), 0, 0, 0])
    for k in sorted(set(ga) - keys):
        ses, pur, rev = ga[k]
        agn, kw = glab[k]
        un.append([to_date(k[0]), k[1], agn, kw, 'GA만',
                   0, 0, 0, 0, 0, ses, pur, rev])
    aonly = [r for r in un if r[4] == '광고만']
    gonly = [r for r in un if r[4] == 'GA만']
    un_show = [r for r in un if r[7] >= a.min_clicks or r[10] > 0]
    un_show.sort(key=lambda r: (-r[7], -r[10]))

    def sm(rs, idx):
        return [sum(r[i] for r in rs) for i in idx]

    b1 = [['구분', '조합 수', '비용', '노출수', '클릭수', '세션수', '구매',
           '구매 수익'],
          ['매칭 (양쪽)', len(rows)] + sm(rows, (4, 5, 6, 9, 10, 11)),
          ['광고만 (GA 세션 없음)', len(aonly)] + sm(aonly, (5, 6, 7, 10, 11, 12)),
          ['GA만 (광고 클릭 없음)', len(gonly)] + sm(gonly, (5, 6, 7, 10, 11, 12)),
          ['합계', len(rows) + len(un),
           sum(r[4] for r in rows) + sum(r[5] for r in un),
           sum(r[5] for r in rows) + sum(r[6] for r in un),
           sum(r[6] for r in rows) + sum(r[7] for r in un),
           sum(r[9] for r in rows) + sum(r[10] for r in un),
           sum(r[10] for r in rows) + sum(r[11] for r in un),
           sum(r[11] for r in rows) + sum(r[12] for r in un)]]

    b2 = [['캠페인', 'GA 추적', '클릭수', '매칭 클릭', '미매칭 클릭', '매칭률']]
    mclk = defaultdict(float)
    for r in rows:
        mclk[r[1]] += r[6]
    for cm, v in sorted(camps.items(), key=lambda x: -x[1][2]):
        m = mclk.get(cm, 0.0)
        b2.append([cm, 'O' if cm in gcamps else 'X', v[2], m, v[2] - m,
                   dv(m, v[2])])
    tc = sum(v[2] for v in camps.values())
    tm = sum(mclk.values())
    b2.append(['합계', '', tc, tm, tc - tm, dv(tm, tc)])

    def why(r):
        if r[1] not in gcamps:
            return 'GA 미추적 캠페인 (PMax 등)'
        return ('클릭 0 (노출만 발생)' if r[7] == 0 else
                '클릭 있으나 GA 세션 없음')

    b3 = [['광고만 — 사유', '조합 수', '노출수', '클릭수', '비용', '클릭 비중']]
    tot_clk = sum(r[7] for r in aonly) or 1
    grp = {}
    for r in aonly:
        grp.setdefault(why(r), []).append(r)
    for lb in sorted(grp, key=lambda x: -sum(r[7] for r in grp[x])):
        v = sm(grp[lb], (6, 7, 5))
        b3.append([lb, len(grp[lb]), v[0], v[1], v[2], v[1] / tot_clk])
    v = sm(aonly, (6, 7, 5))
    b3.append(['소계', len(aonly), v[0], v[1], v[2], 1.0])

    # 검증
    src_clk = sum(v[2] for v in ad.values())
    src_ses = sum(v[0] for v in ga.values())
    got_clk = sum(r[6] for r in rows) + sum(r[7] for r in un)
    got_ses = sum(r[9] for r in rows) + sum(r[10] for r in un)
    if abs(src_clk - got_clk) > 1 or abs(src_ses - got_ses) > 1:
        sys.exit('[중단] 합계가 원본과 맞지 않습니다')
    print('검증  클릭 %s / 세션 %s  → 원본과 일치'
          % (format(got_clk, ',.0f'), format(got_ses, ',.0f')))

    ym = period and re.search(r'(\d{4})년\s*(\d{1,2})월', period)
    tag = '%s-%02d' % (ym.group(1), int(ym.group(2))) if ym else 'out'
    out = a.out or os.path.join(a.dir, 'GSA_GA_매칭_%s.xlsx' % tag)

    wb = openpyxl.Workbook()
    ws = wb.create_sheet('매칭')
    ws.append([c[0] for c in COLS])
    for r in rows:
        ws.append(r)
    style(ws, COLS, len(rows) + 1)

    summary_sheet(wb.create_sheet('미매칭 요약'),
                  [('① 매칭 현황', b1), ('② 캠페인별 매칭률', b2),
                   ('③ 광고만 — 왜 GA 세션이 없나', b3)])

    ws = wb.create_sheet('미매칭 Raw')
    ws.append([c[0] for c in UNCOLS])
    for r in un_show:
        ws.append(r)
    style(ws, UNCOLS, len(un_show) + 1)

    wb.remove(wb['Sheet'])
    try:
        wb.save(out)
    except PermissionError:
        sys.exit('[중단] 파일이 열려 있습니다. 닫고 다시 실행하세요: %s'
                 % os.path.basename(out))
    print('결과  매칭 %d행 / 미매칭 %d조합 (Raw 수록 %d행, 클릭 %d회 이상)'
          % (len(rows), len(un), len(un_show), a.min_clicks))
    print('      %s' % out)


if __name__ == '__main__':
    main()
